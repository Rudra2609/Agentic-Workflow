"""Render the Sprint-2 gold documents and their expected labels (deterministic; safe to re-run).

    cd backend && uv run python ../samples/doc_intel/gold/build_gold.py

Writes ``docs/<id>/...`` (PDF, DOCX, scanned PDF, photo PNG, Markdown, text) and ``expected/<id>.json``.
"""

import io
import json
import random
import sys
from pathlib import Path
from typing import Any

import pypdfium2 as pdfium
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageEnhance, ImageFilter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    ListFlowable,
    ListItem,
    PageTemplate,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import content as c  # noqa: E402

DOCS, EXPECTED = HERE / "docs", HERE / "expected"

# ---------------------------------------------------------------------------------------------------- expected labels


def expected_for(blocks: list[Any], extra_deliverables: list[c.Item] | None = None) -> dict[str, Any]:
    out: dict[str, Any] = {"sections": [], "requirements": [], "constraints": [], "deliverables": [],
                           "evaluation_criteria": [], "total_marks": None}

    def add(text: str, item: c.Item | None) -> None:
        if item is None:
            return
        if item.kind == "R":
            out["requirements"].append({"text": text, "priority": item.priority, "type": item.rtype})
        elif item.kind == "C":
            out["constraints"].append({"text": text, "kind": item.ckind})
        elif item.kind == "D":
            out["deliverables"].append({"name": item.name or text, "format": item.fmt, "text": text})

    for block in blocks:
        kind = block[0]
        if kind in ("h1", "h2"):
            out["sections"].append({"level": 1 if kind == "h1" else 2, "title": block[1]})
        elif kind == "p":
            for sentence, item in block[1]:
                add(sentence, item)
        elif kind in ("ul", "ol"):
            for text, item in block[1]:
                add(text, item)
        elif kind in ("marks", "rubric"):
            for crit in block[1]:
                out["evaluation_criteria"].append({
                    "criterion": crit.criterion, "max_marks": crit.marks,
                    "rubric_levels": [{"label": lab, "marks": m} for lab, _, m in crit.levels]})
        elif kind == "total":
            out["total_marks"] = block[1]
    words: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind in ("h1", "h2", "note"):
            words.append(block[1])
        elif kind == "p":
            words += [s for s, _ in block[1]]
        elif kind in ("ul", "ol"):
            words += [t for t, _ in block[1]]
        elif kind in ("marks", "rubric"):
            words += [f"{cr.criterion} {fmt_marks(cr.marks)}" for cr in block[1]]
        elif kind == "deliverables_table":
            words += [" ".join(r) for r in block[1]]
        elif kind == "total":
            words.append(f"Total: {fmt_marks(block[1])} marks")
    out["reference_text"] = "\n".join(words)
    for item in extra_deliverables or []:
        out["deliverables"].append({"name": item.name, "format": item.fmt, "text": item.name})
    return out


# ---------------------------------------------------------------------------------------------------- PDF


STYLES = getSampleStyleSheet()
H1 = ParagraphStyle("h1", parent=STYLES["Heading1"], fontName="Helvetica-Bold", fontSize=16, spaceAfter=8)
H2 = ParagraphStyle("h2", parent=STYLES["Heading2"], fontName="Helvetica-Bold", fontSize=12.5, spaceBefore=8,
                    spaceAfter=4)
BODY = ParagraphStyle("body", parent=STYLES["BodyText"], fontName="Helvetica", fontSize=10.5, leading=14)
NOTE = ParagraphStyle("note", parent=BODY, fontSize=7.5, leading=9, textColor=colors.HexColor("#444444"))
CELL = ParagraphStyle("cell", parent=BODY, fontSize=9, leading=11)
GRID = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                   ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("BACKGROUND", (0, 0), (-1, 0),
                                                                        colors.HexColor("#e8e8e8"))])


def fmt_marks(m: float | None) -> str:
    return "" if m is None else (str(int(m)) if float(m).is_integer() else str(m))


def pdf_flowables(blocks: list[Any], width: float) -> list[Any]:
    out: list[Any] = []
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            out.append(Paragraph(block[1], H1))
        elif kind == "h2":
            out.append(Paragraph(block[1], H2))
        elif kind == "p":
            out.append(Paragraph(" ".join(s for s, _ in block[1]), BODY))
            out.append(Spacer(1, 4))
        elif kind in ("ul", "ol"):
            items = [ListItem(Paragraph(t, BODY), leftIndent=14) for t, _ in block[1]]
            out.append(ListFlowable(items, bulletType="1" if kind == "ol" else "bullet", start=1 if kind == "ol"
                                    else None, leftIndent=14, bulletFontSize=9))
            out.append(Spacer(1, 4))
        elif kind == "marks":
            rows = [["Criterion", "Marks"]] + [[Paragraph(cr.criterion, CELL), fmt_marks(cr.marks)]
                                               for cr in block[1]]
            out.append(Table(rows, colWidths=[width * 0.75, width * 0.2], style=GRID))
            out.append(Spacer(1, 6))
        elif kind == "rubric":
            crits: list[c.Crit] = block[1]
            labels = [lab for lab, _, _ in crits[0].levels]
            rows = [["Section", "Criterion", "Marks", *labels]]
            style = TableStyle(list(GRID.getCommands()))
            start = 1
            for i, cr in enumerate(crits, start=1):
                rows.append([cr.group, Paragraph(cr.criterion, CELL), fmt_marks(cr.marks),
                             *[Paragraph(f"{d} ({fmt_marks(m)})", CELL) for _, d, m in cr.levels]])
                last = i == len(crits) or crits[i].group != cr.group
                if last:
                    if i > start:
                        style.add("SPAN", (0, start), (0, i))
                        for r in range(start + 1, i + 1):
                            rows[r][0] = ""
                    start = i + 1
            n = len(labels)
            widths = [width * 0.17, width * 0.22, width * 0.08] + [width * 0.53 / n] * n
            out.append(Table(rows, colWidths=widths, style=style, repeatRows=1))
            out.append(Spacer(1, 6))
        elif kind == "deliverables_table":
            rows = [list(r) for r in block[1]]
            out.append(Table(rows, colWidths=[width * 0.45, width * 0.35, width * 0.15], style=GRID))
        elif kind == "total":
            out.append(Paragraph(f"<b>Total: {fmt_marks(block[1])} marks</b>", BODY))
        elif kind == "note":
            out.append(Spacer(1, 10))
            out.append(Paragraph(block[1], NOTE))
    return out


def write_pdf(blocks: list[Any], path: Path, *, two_column: bool = False, header: str = "") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    margin = 2 * cm
    width = A4[0] - 2 * margin

    def decorate(canvas: Any, doc: Any) -> None:
        if header:
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.drawString(margin, A4[1] - 1.2 * cm, header)
            canvas.drawRightString(A4[0] - margin, 1.2 * cm, f"Page {doc.page}")
            canvas.restoreState()

    if two_column:
        doc = BaseDocTemplate(str(path), pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=margin,
                              bottomMargin=margin, title=path.stem, author="IT314 Group 23", invariant=1)
        gap = 0.8 * cm
        col = (width - gap) / 2
        frames = [Frame(margin, margin, col, A4[1] - 2 * margin, id="left"),
                  Frame(margin + col + gap, margin, col, A4[1] - 2 * margin, id="right")]
        doc.addPageTemplates([PageTemplate(id="two", frames=frames, onPage=decorate)])
        doc.build(pdf_flowables(blocks, col))
        return
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=margin,
                            bottomMargin=margin, title=path.stem, author="IT314 Group 23", invariant=1)
    doc.build(pdf_flowables(blocks, width), onFirstPage=decorate, onLaterPages=decorate)


# ---------------------------------------------------------------------------------------------------- DOCX


def write_docx(blocks: list[Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.author = "IT314 Group 23"
    doc.styles["Normal"].font.size = Pt(11)
    for block in blocks:
        kind = block[0]
        if kind in ("h1", "h2"):
            doc.add_heading(block[1], level=1 if kind == "h1" else 2)
        elif kind == "p":
            doc.add_paragraph(" ".join(s for s, _ in block[1]))
        elif kind in ("ul", "ol"):
            for text, _ in block[1]:
                doc.add_paragraph(text, style="List Number" if kind == "ol" else "List Bullet")
        elif kind in ("marks", "rubric", "deliverables_table"):
            if kind == "marks":
                rows = [["Criterion", "Marks"]] + [[cr.criterion, fmt_marks(cr.marks)] for cr in block[1]]
            elif kind == "deliverables_table":
                rows = [list(r) for r in block[1]]
            else:
                labels = [lab for lab, _, _ in block[1][0].levels]
                rows = [["Section", "Criterion", "Marks", *labels]] + [
                    [cr.group, cr.criterion, fmt_marks(cr.marks), *[f"{d} ({fmt_marks(m)})" for _, d, m in cr.levels]]
                    for cr in block[1]]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for col, value in enumerate(row):
                    table.cell(r, col).text = value
            if kind == "rubric":  # merge the Section column over each group
                start = 1
                for r in range(1, len(rows) + 1):
                    if r == len(rows) or rows[r][0] != rows[start][0]:
                        if r - 1 > start:
                            merged = table.cell(start, 0).merge(table.cell(r - 1, 0))
                            merged.text = rows[start][0]
                        start = r
        elif kind == "total":
            doc.add_paragraph().add_run(f"Total: {fmt_marks(block[1])} marks").bold = True
        elif kind == "note":
            doc.add_paragraph(block[1]).runs[0].font.size = Pt(8)
    # python-docx stamps the current time; fix it so re-runs give identical files
    from datetime import datetime
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 1)
    doc.save(str(path))


# ---------------------------------------------------------------------------------------------------- scans


def render_pages(pdf_path: Path, dpi: int) -> list[Image.Image]:
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        return [page.render(scale=dpi / 72).to_pil().convert("L") for page in pdf]
    finally:
        pdf.close()


def scanify(img: Image.Image, rng: random.Random, angle: float) -> Image.Image:
    img = img.rotate(angle, resample=Image.Resampling.BICUBIC, expand=True, fillcolor=255)
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    px = img.load()
    w, h = img.size
    for _ in range(w * h // 400):  # speckle noise
        x, y = rng.randrange(w), rng.randrange(h)
        px[x, y] = rng.choice((0, 90, 200))
    return ImageEnhance.Contrast(img).enhance(0.85)


def write_scanned_pdf(blocks: list[Any], path: Path) -> None:
    tmp = path.with_suffix(".src.pdf")
    write_pdf(blocks, tmp)
    rng = random.Random(4)
    pages = [scanify(p, rng, angle=1.2) for p in render_pages(tmp, dpi=150)]
    tmp.unlink()
    pages[0].save(str(path), "PDF", resolution=150, save_all=True, append_images=pages[1:])


def write_photo(blocks: list[Any], path: Path) -> None:
    tmp = path.with_suffix(".src.pdf")
    write_pdf(blocks, tmp)
    page = render_pages(tmp, dpi=130)[0]
    tmp.unlink()
    page = page.crop((0, 0, page.width, int(page.height * 0.42)))
    img = scanify(page, random.Random(5), angle=-2.5).convert("RGB")
    tint = Image.new("RGB", img.size, (255, 236, 205))
    img = Image.blend(img, tint, 0.18)
    img = ImageEnhance.Brightness(img).enhance(0.93)
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(buf.getvalue())


# ---------------------------------------------------------------------------------------------------- text


def write_text(blocks: list[Any], path: Path, *, markdown: bool) -> None:
    lines: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind in ("h1", "h2"):
            lines += [("# " if kind == "h1" else "## ") + block[1] if markdown else block[1], ""]
        elif kind == "p":
            lines += [" ".join(s for s, _ in block[1]), ""]
        elif kind in ("ul", "ol"):
            lines += [(f"{i}. " if kind == "ol" else "- ") + t for i, (t, _) in enumerate(block[1], start=1)]
            lines.append("")
        elif kind == "marks":
            lines += ["| Criterion | Marks |", "|---|---|"]
            lines += [f"| {cr.criterion} | {fmt_marks(cr.marks)} |" for cr in block[1]]
            lines.append("")
        elif kind == "total":
            lines += [f"Total: {fmt_marks(block[1])} marks", ""]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------------------------------- the set

GOLD: list[dict[str, Any]] = [
    {"id": "g01_tcp_chat", "title": "Programming assignment (native PDF)",
     "files": [("tcp_chat_assignment.pdf", None, c.TCP_CHAT, "pdf")]},
    {"id": "g02_lab_manual", "title": "Lab manual with numbered steps (DOCX)",
     "files": [("lab4_unit_testing.docx", None, c.LAB_MANUAL, "docx")]},
    {"id": "g03_rubric", "title": "Rubric-heavy marking scheme, merged cells (PDF)",
     "files": [("marking_scheme.pdf", None, c.RUBRIC, "pdf")]},
    {"id": "g04_scanned", "title": "Scanned assignment (image-only PDF, skewed)", "ocr": True,
     "files": [("library_assignment_scan.pdf", None, c.SCANNED, "scan")]},
    {"id": "g05_photo", "title": "Phone photo of a handout (PNG)", "ocr": True,
     "files": [("quiz_handout_photo.png", None, c.PHOTO, "photo")]},
    {"id": "g06_project_brief", "title": "Client brief with styles, lists and a deliverables table (DOCX)",
     "files": [("event_booking_brief.docx", None, c.PROJECT_BRIEF, "docx")],
     "extra_deliverables": c.PROJECT_BRIEF_DELIVERABLES},
    {"id": "g07_messy", "title": "Two-column layout with header, footer and a footnote (PDF)",
     "files": [("weather_dashboard.pdf", None, c.MESSY, "pdf2col")]},
    {"id": "g08_multipart", "title": "Multi-part assignment uploaded as folders",
     "files": [("part_a.pdf", "Part A", c.PART_A, "pdf"), ("part_b.docx", "Part B", c.PART_B, "docx"),
               ("part_c.md", "Part C", c.PART_C, "md")]},
    {"id": "g09_request", "title": "Typed request (no upload)",
     "files": [("request.txt", None, c.REQUEST, "txt")]},
    {"id": "g10_long_lab", "title": "Longer lab sheet with requirements in lists under lead-in sentences (PDF)",
     "files": [("parking_garage_lab.pdf", None, c.LONG_LAB, "pdf")]},
]


def build() -> None:
    EXPECTED.mkdir(parents=True, exist_ok=True)
    for doc in GOLD:
        files_out = []
        merged: dict[str, Any] = {}
        for name, folder, blocks, kind in doc["files"]:
            rel = Path(doc["id"]) / (folder or "") / name
            target = DOCS / rel
            if kind == "pdf":
                write_pdf(blocks, target)
            elif kind == "pdf2col":
                write_pdf(blocks, target, two_column=True, header="IT 314 Software Engineering - Mini project brief")
            elif kind == "docx":
                write_docx(blocks, target)
            elif kind == "scan":
                write_scanned_pdf(blocks, target)
            elif kind == "photo":
                write_photo(blocks, target)
            elif kind in ("md", "txt"):
                write_text(blocks, target, markdown=kind == "md")
            exp = expected_for(blocks, doc.get("extra_deliverables") if len(doc["files"]) == 1 else None)
            for key in ("requirements", "constraints", "deliverables", "evaluation_criteria"):
                for entry in exp[key]:
                    entry["file"] = name
                    entry["folder"] = folder
            files_out.append({"path": str(rel), "name": name, "folder": folder, "sections": exp["sections"],
                              "total_marks": exp["total_marks"], "reference_text": exp["reference_text"]})
            for key in ("requirements", "constraints", "deliverables", "evaluation_criteria"):
                merged.setdefault(key, []).extend(exp[key])
        totals = [f["total_marks"] for f in files_out if f["total_marks"] is not None]
        record = {"id": doc["id"], "title": doc["title"], "needs_ocr": bool(doc.get("ocr")), "files": files_out,
                  "total_marks": sum(totals) if totals else None, **merged}
        (EXPECTED / f"{doc['id']}.json").write_text(json.dumps(record, indent=2) + "\n", encoding="utf-8")
        counts = {k: len(merged[k]) for k in ("requirements", "constraints", "deliverables", "evaluation_criteria")}
        print(f"{doc['id']:20s} {len(files_out)} file(s)  {counts}")


if __name__ == "__main__":
    build()
